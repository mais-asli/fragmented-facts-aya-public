"""Merge language-owned review cells without assigning missing judgments.

The returned partner workbooks are read-only inputs. This writes a structured
intermediate for the final review workbook and a field-level provenance log.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data/review/returned_20260925"
OUT = ROOT / "data/review/merged_20260925"
FILES = {
    "he": SRC / "Aya_Pilot_Review_Hebrew_20260925.xlsx",
    "ar": SRC / "Aya_Pilot_Review_Arabic_20260920.xlsx",
}
SCHEMA = json.loads((ROOT / "outputs/01a09b51/pilot_review/workbook_schema.json").read_text(encoding="utf-8"))
FACT_RELATIONS = {
    "Place of birth": "P19",
    "Place of death": "P20",
    "Headquarters location": "P159",
    "Location of formation": "P740",
}


def read(path: Path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        result = {}
        for name, spec in SCHEMA.items():
            sheet = book[name]
            rows = sheet.iter_rows(min_row=7, max_row=7 + spec["row_count"],
                                   max_col=len(spec["headers"]), values_only=True)
            headers = list(next(rows))
            if headers != spec["headers"]:
                raise ValueError(f"Unexpected headers in {path}: {name}")
            result[name] = [dict(zip(headers, values)) for values in rows]
        return result
    finally:
        book.close()


def key(sheet, row):
    return row[{"Pilot facts": "Fact ID", "Hebrew": "Fact ID", "Arabic": "Fact ID",
                "Templates": "Template key", "Smoke review": "Example"}[sheet]]


def main():
    if OUT.exists():
        raise FileExistsError(OUT)
    books = {language: read(path) for language, path in FILES.items()}
    base = read(ROOT / "outputs/01a09b51/pilot_review/Aya_Pilot_Review.xlsx")
    merged = {sheet: [] for sheet in SCHEMA}
    changes = []
    for sheet in SCHEMA:
        for original in base[sheet]:
            identity = key(sheet, original)
            versions = {lang: next(row for row in books[lang][sheet] if key(sheet, row) == identity)
                        for lang in books}
            item = dict(original)
            if sheet == "Pilot facts":
                relation = FACT_RELATIONS[item["Relation"]]
                owner = "he" if relation in ("P19", "P20") else "ar"
                # Keep the six already-source-checked reserve P19/P20 rows from
                # the Arabic workbook when the Hebrew workbook has no evidence.
                if not versions[owner]["Evidence URL"] and versions["ar"]["Evidence URL"]:
                    owner = "ar"
                for field in ("Decision", "Reviewer", "Before May 2024", "Evidence URL", "Review notes"):
                    item[field] = versions[owner][field]
                for field in ("Subject in Hebrew", "Answer in Hebrew", "Accepted Hebrew names (JSON)"):
                    item[field] = versions["he"][field]
                for field in ("Subject in Arabic", "Answer in Arabic", "Accepted Arabic names (JSON)"):
                    item[field] = versions["ar"][field]
                # Never replace an evidence URL with the generic Wikidata entity
                # link. Keep it separately in the original Source link column.
            elif sheet in ("Hebrew", "Arabic"):
                item = dict(versions["he" if sheet == "Hebrew" else "ar"])
            elif sheet == "Templates":
                language = item["Language"]
                owner = ("he" if language == "he" else "ar" if language == "ar"
                         else "he" if item["Relation"] in ("Place of birth", "Place of death") else "ar")
                item = dict(versions[owner])
            elif sheet == "Smoke review":
                item = dict(versions["he" if identity in (1, 3, 4) else "ar"])
            for field, new_value in item.items():
                if new_value != original[field]:
                    changes.append({"sheet": sheet, "id": identity, "field": field,
                                    "original": original[field], "merged": new_value})
            merged[sheet].append(item)
    facts = merged["Pilot facts"]
    no_evidence = [{"fact_id": r["Fact ID"], "subject": r["Subject in English"],
                    "relation": r["Relation"]} for r in facts if not r["Evidence URL"]]
    OUT.mkdir(parents=True)
    (OUT / "merged_input.json").write_text(json.dumps(merged, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "field_changes.json").write_text(json.dumps(changes, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "summary.json").write_text(json.dumps({
        "status": "merged_review_inputs_not_final_approval",
        "facts_with_direct_evidence_url": len(facts) - len(no_evidence),
        "facts_missing_direct_evidence_url": no_evidence,
        "hebrew_pairs_with_marked_spelling": sum(bool(r["Marked spelling"]) for r in merged["Hebrew"]),
        "arabic_pairs_with_marked_spelling": sum(bool(r["Marked spelling"]) for r in merged["Arabic"]),
        "merged_field_changes_from_original": len(changes),
        "source_files": {k: str(v) for k, v in FILES.items()},
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print((OUT / "summary.json").read_text(encoding="utf-8"))


if __name__ == "__main__":
    main()
