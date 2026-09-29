"""Extract ReviewerA's saved Review sheet to a UTF-8 CSV for structural checks.

This only reads the workbook. It does not infer human approval from the AI
suggestion columns and it does not overwrite the original blank review queue.
"""
import argparse
import csv
import json
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
BATCH = ROOT / "data" / "review" / "heldout_batch_v1"
PAYLOAD = ROOT / "outputs" / "01a09b51" / "heldout_review_workbook_input.json"
DEFAULT_BOOK = ROOT / "outputs" / "01a09b51" / "Aya_Independent_Test_Review_20260927.xlsx"
DEFAULT_CSV = BATCH / "AYA_ONLY_TEST_REVIEW_FROM_WORKBOOK_20260927.csv"

parser = argparse.ArgumentParser()
parser.add_argument("--workbook", type=Path, default=DEFAULT_BOOK)
parser.add_argument("--output", type=Path, default=DEFAULT_CSV)
args = parser.parse_args()

columns = json.loads(PAYLOAD.read_text(encoding="utf-8"))["columns"]
labels = [label for _, label in columns]
book = load_workbook(args.workbook, read_only=True, data_only=True)
sheet = book["Review"]
sheet_rows = list(sheet.iter_rows(min_row=1, max_row=61, min_col=1,
                                 max_col=len(columns), values_only=True))
actual = list(sheet_rows[0])
if actual != labels:
    raise ValueError("Review headers/order changed; restore the original workbook structure")

original_header = list(csv.DictReader((BATCH / "AYA_ONLY_TEST_REVIEW_QUEUE_20260927.csv")
                                      .open(encoding="utf-8-sig", newline="")).fieldnames)
args.output.parent.mkdir(parents=True, exist_ok=True)
with args.output.open("w", encoding="utf-8-sig", newline="") as fh:
    writer = csv.DictWriter(fh, fieldnames=original_header)
    writer.writeheader()
    for excel_row, values in enumerate(sheet_rows[1:], 2):
        item = {}
        for (key, _), value in zip(columns, values):
            if isinstance(value, (date, datetime)):
                value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            item[key] = "" if value is None else str(value)
        writer.writerow({key: item.get(key, "") for key in original_header})
print(args.output)
