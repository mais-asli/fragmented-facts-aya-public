"""Read the reviewed Excel workbook and export validated copies without overwriting reviews.

Uses openpyxl only for reading. Run with the bundled workspace Python runtime.
"""
import argparse
import copy
from datetime import datetime, timezone
import json
from pathlib import Path
import sys
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from fragmented_facts.data import csv_rows, write_csv
from fragmented_facts.io import digest, file_hash, read_json, write_json
from fragmented_facts.prompts import validate_templates
from fragmented_facts.unicode import validate_pair


def text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value).lower()
    return str(value).strip()


def read_rows(book, schema, name):
    sheet = book[name]
    spec = schema[name]
    headers = [text(sheet.cell(7, c + 1).value) for c in range(len(spec['headers']))]
    if headers != spec['headers']:
        raise ValueError('Workbook headers changed: ' + name)
    if any(cell.value is not None for row in sheet.iter_rows(min_row=8+spec['row_count']) for cell in row):
        raise ValueError('Additional rows need a separate review batch: ' + name)
    return [{header: text(sheet.cell(r, c + 1).value) for c, header in enumerate(headers)}
            for r in range(8, 8 + spec['row_count'])]


def merge(current, baseline, incoming, key):
    live = {key(row): row for row in current}
    original = {key(row): row for row in baseline}
    for row in incoming:
        identity = key(row)
        if identity not in live or identity not in original:
            raise ValueError('Unknown review identity')
        if row == original[identity]:
            continue
        if live[identity] != original[identity] and live[identity] != row:
            raise ValueError('Another review changed the same record; reconcile: ' + str(identity))
        live[identity] = row
    return [live[key(row)] for row in current]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', default='outputs/01a09b51/pilot_review/Aya_Pilot_Review.xlsx')
    parser.add_argument('--check-only', action='store_true')
    parser.add_argument('--output')
    args = parser.parse_args()
    from openpyxl import load_workbook
    packet = read_json(ROOT / 'data/review/pilot_batch_v1/workbook_input.json')
    manifest = copy.deepcopy(packet['manifest'])
    expected = manifest.pop('manifest_hash')
    if digest(manifest) != expected:
        raise ValueError('Pilot review manifest was edited')
    if file_hash(ROOT / 'data/review/subject_splits.json') != manifest['split_manifest_sha256']:
        raise ValueError('Pinned split assignments changed')
    workbook_path = ROOT / args.workbook
    schema = read_json(ROOT / 'outputs/01a09b51/pilot_review/workbook_schema.json')
    book = load_workbook(workbook_path, read_only=False, data_only=True)
    try:
        if book.sheetnames[:len(schema)] != list(schema):
            raise ValueError('Required review tabs changed or reordered')
        facts = read_rows(book, schema, 'Pilot facts')
        expected_ids = {f['fact_id'] for f in manifest['facts']}
        if len(facts) != len(expected_ids) or {f['Fact ID'] for f in facts} != expected_ids:
            raise ValueError('Fact IDs changed or duplicated')
        original = {f['fact_id']: f for f in packet['facts']}
        facts_out = []
        fields = {'subject_en':'Subject in English', 'object_en':'Answer in English',
                  'subject_he':'Subject in Hebrew', 'subject_ar':'Subject in Arabic',
                  'object_he':'Answer in Hebrew', 'object_ar':'Answer in Arabic',
                  'aliases_en':'Accepted English names (JSON)', 'aliases_he':'Accepted Hebrew names (JSON)',
                  'aliases_ar':'Accepted Arabic names (JSON)', 'decision':'Decision', 'reviewer':'Reviewer',
                  'pre_release_verified':'Before May 2024', 'evidence_url':'Evidence URL', 'notes':'Review notes'}
        for row in facts:
            base = dict(original[row['Fact ID']])
            if row['Relation'] != packet['relation_names'][base['relation']]:
                raise ValueError('Relation identity changed in review workbook')
            if row['Subject QID'] != base['subject_qid'] or row['Answer QID'] != base['object_qid']:
                raise ValueError('Entity identity changed in review workbook')
            for key, label in fields.items():
                base[key] = row[label]
            if base['decision'] not in ('pending','approved','rejected'):
                raise ValueError('Invalid fact decision')
            for language in ('en','he','ar'):
                aliases = json.loads(base['aliases_' + language])
                if not isinstance(aliases,list) or not aliases or not all(isinstance(v,str) and v.strip() for v in aliases):
                    raise ValueError('Accepted names must be a nonempty JSON list')
            if base['decision'] == 'approved':
                if not base['reviewer'] or base['pre_release_verified'].lower() != 'true' or urlparse(base['evidence_url']).scheme not in ('http','https'):
                    raise ValueError('Approved fact lacks reviewer, date check or evidence URL')
                if any(not base[key] for key in fields if key.startswith(('subject_','object_'))):
                    raise ValueError('Approved fact is missing a canonical name')
            facts_out.append(base)
        reviewed_facts = {f['fact_id']: f for f in facts_out}
        original_pairs = {(r['fact_id'],r['language']):r for r in packet['pairs']}
        pairs_out = []
        for sheet, language in (('Hebrew','he'),('Arabic','ar')):
            rows = read_rows(book,schema,sheet)
            if len(rows) != len(expected_ids) or {r['Fact ID'] for r in rows} != expected_ids:
                raise ValueError('Pair IDs changed or duplicated')
            for row in rows:
                item = dict(original_pairs[row['Fact ID'],language])
                item.update(U=row['Unmarked spelling'], D=row['Marked spelling'], decision=row['Decision'],
                            reviewer=row['Reviewer'], notes=row['Review notes'])
                if row['Spelling source URL']:
                    item['notes'] += ('\n' if item['notes'] else '') + 'Spelling source: ' + row['Spelling source URL']
                if item['decision'] not in ('pending','approved','rejected'):
                    raise ValueError('Invalid pair decision')
                if item['decision'] == 'approved':
                    validate_pair(item['U'],item['D'],language)
                    if not item['reviewer']:
                        raise ValueError('Approved spelling needs a named reviewer')
                    if item['U'] != reviewed_facts[item['fact_id']]['subject_'+language]:
                        raise ValueError('Unmarked spelling differs from canonical fact label; recalculate workbook')
                pairs_out.append(item)
        templates = read_json(ROOT / 'configs/templates.json')
        rows = read_rows(book,schema,'Templates')
        if len(rows) != len(manifest['template_keys']) or {r['Template key'] for r in rows} != set(manifest['template_keys']):
            raise ValueError('Template keys changed')
        for row in rows:
            stage, relation, language, template_id = row['Template key'].split('/')
            if row['Decision'] not in ('pending','approved','rejected'):
                raise ValueError('Invalid template decision')
            if row['Decision'] == 'approved' and not row['Reviewer']:
                raise ValueError('Approved template needs reviewer')
            item = templates[stage][relation][language][template_id]
            item.update(text=row['Question text'], status='pending_native_review' if row['Decision']=='pending' else row['Decision'],
                        reviewer=row['Reviewer'], notes=row['Review notes'])
        validate_templates(templates, require_review=False)
        smoke_rows = read_rows(book,schema,'Smoke review')
        for i,row in enumerate(smoke_rows):
            expected_example = packet['smoke'][i]
            if (row['Subject shown to Aya'] != expected_example['subject'] or row['Aya answer'] != expected_example['answer']
                    or row['Language'] != expected_example['language']):
                raise ValueError('Original smoke evidence changed')
            if row['Review state'] not in ('pending','reviewed'):
                raise ValueError('Invalid smoke review state')
            if row['Review state']=='reviewed' and (not row['Reviewer'] or row['Spelling valid'] not in ('yes','no','unsure')
                                                   or row['Answer correct'] not in ('yes','no','unsure')):
                raise ValueError('Smoke review has missing judgments or reviewer')
        config = read_json(ROOT / 'configs/study.json')
        if all(r['Review state']=='reviewed' for r in smoke_rows):
            config['native_smoke_review'] = {'reviewer': ', '.join(sorted({r['Reviewer'] for r in smoke_rows})),
                                            'notes': 'See smoke_review.json exported from the reviewed pilot workbook.'}
        complete = [f for f in facts_out if f['decision']=='approved'
                    and all(any(p['fact_id']==f['fact_id'] and p['language']==l and p['decision']=='approved' for p in pairs_out) for l in ('he','ar'))]
        summary = {'status':'review_inputs_validated', 'approved_facts':sum(f['decision']=='approved' for f in facts_out),
                   'approved_pairs':sum(p['decision']=='approved' for p in pairs_out),
                   'subjects_with_both_languages_reviewed':len(complete), 'required_pilot_subjects':50,
                   'approved_templates':sum(r['Decision']=='approved' for r in rows),
                   'required_templates':len(rows), 'smoke_reviews':sum(r['Review state']=='reviewed' for r in smoke_rows),
                   'workbook_sha256':file_hash(workbook_path), 'check_only':args.check_only}
        write_json(ROOT / 'results/pilot_review_readiness.json', summary)
        if not args.check_only:
            out = ROOT / (args.output or ('data/review/pilot_batch_v1/imported_' + datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')))
            if out.exists():
                raise FileExistsError('Review output exists; preserve previous import')
            merged_facts=merge(csv_rows(ROOT / 'data/review/facts.csv'),packet['facts'],facts_out,lambda r:r['fact_id'])
            merged_pairs=merge(csv_rows(ROOT / 'data/review/pairs.csv'),packet['pairs'],pairs_out,lambda r:(r['fact_id'],r['language']))
            out.mkdir(parents=True)
            write_csv(out / 'facts.csv',merged_facts,list(merged_facts[0]))
            write_csv(out / 'pairs.csv',merged_pairs,list(merged_pairs[0]))
            write_json(out / 'templates.json',templates)
            write_json(out / 'study.json',config)
            write_json(out / 'smoke_review.json',smoke_rows)
            write_json(out / 'summary.json',summary)
            summary['output']=str(out)
        print(json.dumps(summary,indent=2))
    finally:
        book.close()

if __name__=='__main__':
    main()
